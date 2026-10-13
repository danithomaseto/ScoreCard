"""Datas do export em dia/mes (Brasil) e mes/dia (servidor americano).

O caso real: na Hughes a semana veio como "9/6/26" (6 de setembro, mes
primeiro). Lida como dia/mes, virava 9 de junho, o "10/4/26" virava 10 de
abril — e Junho e Abril apareciam no filtro de mes — e o "9/13/26" nao
virava data nenhuma: a semana sumia do calculo.
"""

import datetime

import pytest

import coverage_store
import grupos_store
import indicators_store
import indiretas_store
import reparo_semanas
from indicators import datas, pico, reader

HOJE = datetime.date(2026, 10, 13)


def semanas(valores):
    return list(datas.semanas(valores, hoje=HOJE).values())


# ---------------- Ordem da coluna ----------------

def test_semana_americana_curta_nao_troca_dia_e_mes_nem_some():
    assert semanas(["8/30/26", "9/6/26", "9/13/26", "9/20/26", "9/27/26", "10/4/26"]) == [
        "2026-08-30", "2026-09-06", "2026-09-13", "2026-09-20", "2026-09-27", "2026-10-04"]


def test_semana_brasileira_continua_como_era():
    assert semanas(["30/08/2026", "06/09/2026", "13/09/2026"]) == [
        "2026-08-30", "2026-09-06", "2026-09-13"]


@pytest.mark.parametrize("valores, esperado", [
    # So datas ambiguas (dia e mes ate 12): decide o dia da semana.
    (["9/6/26", "10/4/26"], ["2026-09-06", "2026-10-04"]),        # americana, domingo
    (["06/09/2026", "04/10/2026"], ["2026-09-06", "2026-10-04"]),  # brasileira, domingo
    (["9/6/2026", "10/4/2026"], ["2026-09-06", "2026-10-04"]),    # americana com ano inteiro
    (["10/4/26"], ["2026-10-04"]),                                 # uma so
    (["04/10/2026"], ["2026-10-04"]),
    (["9/7/26", "9/14/26"], ["2026-09-07", "2026-09-14"]),        # semana que comeca na segunda
    (["07/09/2026"], ["2026-09-07"]),
])
def test_data_ambigua_vai_para_o_dia_em_que_a_semana_comeca(valores, esperado):
    assert semanas(valores) == esperado


def test_celula_de_data_e_iso_entram_como_vieram():
    assert semanas([datetime.datetime(2026, 9, 27), datetime.datetime(2026, 10, 4)]) == [
        "2026-09-27", "2026-10-04"]
    assert semanas(["2026-09-06", "2026-09-13 00:00:00"]) == ["2026-09-06", "2026-09-13"]


def test_datas_certas_da_coluna_ensinam_o_dia_da_semana():
    # Celulas de data em domingo; o texto ambiguo segue o domingo.
    lidas = datas.semanas([datetime.datetime(2026, 9, 27), "10/4/26"], hoje=HOJE)
    assert lidas[datas.chave("10/4/26")] == "2026-10-04"


def test_nome_de_gestor_e_turno_nao_sao_data():
    assert not datas.e_data("SILVA,ANA")
    assert not datas.e_data("T1")
    assert not datas.e_data(46292.0)
    assert datas.e_data("9/13/26")


# ---------------- Leitor ----------------

TITULOS = ["High Level\ngroup", "Medium Level\nGroup", "Detail Level\nGroup", "Var", "Goal",
           "Measured\nDirect", "Unmeasured\nSignon Direct", "Unmeasured\nInsert Direct", "PD Brk", "Total"]


def _export(tmp_path, semanas_do_arquivo, formato_de_celula=None):
    from openpyxl import Workbook

    caminho = tmp_path / "summary.xlsx"
    livro = Workbook()
    aba = livro.active
    aba.append(TITULOS)
    for i, semana in enumerate(semanas_do_arquivo):
        for gestor in ("SILVA,ANA", "SOUZA,BRUNO"):
            aba.append([semana, gestor, f"P{i:02d}", 0, 30.0, 30.0, 1.0, 0.0, 1.0, 40.0])
            if formato_de_celula:
                aba.cell(aba.max_row, 1).number_format = formato_de_celula
    livro.save(caminho)
    return str(caminho)


def test_leitor_com_semana_americana_em_texto(tmp_path):
    lido = reader.ler(_export(tmp_path, ["8/30/26", "9/6/26", "9/13/26", "10/4/26"]))

    assert lido["tem_semana"] and lido["tem_grupo"]
    assert sorted({linha["semana"] for linha in lido["linhas"]}) == [
        "2026-08-30", "2026-09-06", "2026-09-13", "2026-10-04"]
    assert all(linha["semana"] for linha in lido["linhas"])  # nenhuma linha perdeu a semana


def test_leitor_com_celula_de_data_no_formato_americano(tmp_path):
    # Como veio o arquivo da Hughes: data de verdade, exibida m/d/yy.
    celulas = [datetime.datetime(2026, 9, 27), datetime.datetime(2026, 10, 4)]
    lido = reader.ler(_export(tmp_path, celulas, formato_de_celula="m/d/yy"))

    assert sorted({linha["semana"] for linha in lido["linhas"]}) == ["2026-09-27", "2026-10-04"]


def test_semana_americana_no_medium_continua_sendo_semana(tmp_path):
    """Week > User ID: o Medium so e semana se todos os valores forem
    data — "8/30/26" tambem precisa contar como data."""
    from openpyxl import Workbook

    caminho = tmp_path / "dois_niveis.xlsx"
    livro = Workbook()
    aba = livro.active
    aba.append(TITULOS[1:])
    for semana in ("8/30/26", "9/6/26"):
        aba.append([semana, "P01", 0, 30.0, 30.0, 1.0, 0.0, 1.0, 40.0])
    livro.save(caminho)

    lido = reader.ler(str(caminho))
    assert lido["tem_semana"] and not lido["tem_grupo"]
    assert [linha["semana"] for linha in lido["linhas"]] == ["2026-08-30", "2026-09-06"]


# ---------------- Dias de pico ----------------

def _dia(data, diretas):
    return {"detalhe": data, "measured_direct": diretas, "signon_direct": 0.0,
            "insert_direct": 0.0, "total": 100.0, "pd_brk": 0.0}


def test_pico_com_report_date_americano():
    linhas = [_dia("9/1/26", 90), _dia("9/2/26", 91), _dia("9/3/26", 92), _dia("9/4/26", 93),
              _dia("9/8/26", 94), _dia("9/9/26", 95)]
    resultado = pico.calcular(linhas, "2026-09")

    assert [d["data"] for d in resultado["dias"]] == [
        "2026-09-09", "2026-09-08", "2026-09-04", "2026-09-03", "2026-09-02"]
    assert resultado["ignorados"] == []


def test_pico_com_report_date_brasileiro_e_celula():
    linhas = [_dia("01/09/2026", 90), _dia("02/09/2026", 91), _dia("2026-09-03 00:00:00", 92)]
    resultado = pico.calcular(linhas, "2026-09")

    assert sorted(d["data"] for d in resultado["dias"]) == ["2026-09-01", "2026-09-02", "2026-09-03"]


# ---------------- Reparo do que ja estava gravado ----------------

def test_reparo_acerta_semanas_gravadas_com_dia_e_mes_trocados():
    # Como ficou a Hughes: 06/09 gravada como 09/06 e 04/10 como 10/04,
    # ao lado de uma semana certa (27/09, domingo).
    indicators_store.salvar_extracao("hughes", "week", {
        "2026-06-09": {"efetividade": 0.95}, "2026-04-10": {"efetividade": 0.97},
        "2026-09-27": {"efetividade": 0.99}})
    indicators_store.salvar_extracao("hughes", "month", {"2026-09": {"efetividade": 0.98}})
    grupos_store.salvar("gestor", "hughes", "week", {"2026-06-09": {"SILVA,ANA": {}}})
    coverage_store.salvar_extracao("hughes", "week", {"2026-04-10": {"P01": {"lms": 40}}})
    indiretas_store.salvar("hughes", {"2026-06-09": {"horas_totais": 100}})

    assert reparo_semanas.reparar() == 5

    semanas_gravadas = indicators_store.get_indicators("hughes")["week"]
    assert sorted(semanas_gravadas) == ["2026-09-06", "2026-09-27", "2026-10-04"]
    assert semanas_gravadas["2026-09-06"]["efetividade"] == 0.95
    assert list(indicators_store.get_indicators("hughes")["month"]) == ["2026-09"]
    assert list(grupos_store.da_operacao("gestor", "hughes")) == ["week:2026-09-06"]
    assert list(coverage_store.extracoes("hughes")["hughes"]) == ["week:2026-10-04"]
    assert list(indiretas_store.da_operacao("hughes")) == ["2026-09-06"]

    assert reparo_semanas.reparar() == 0  # nada mais a fazer: nao regrava


def test_reparo_fica_com_a_semana_certa_quando_ela_ja_foi_extraida_de_novo():
    indicators_store.salvar_extracao("hughes", "week", {"2026-06-09": {"efetividade": 0.5}})
    indicators_store.salvar_extracao("hughes", "week", {"2026-09-06": {"efetividade": 0.9}})

    reparo_semanas.reparar()

    semanas_gravadas = indicators_store.get_indicators("hughes")["week"]
    assert list(semanas_gravadas) == ["2026-09-06"]
    assert semanas_gravadas["2026-09-06"]["efetividade"] == 0.9


def test_reparo_nao_mexe_em_semana_de_domingo_ou_segunda_nem_no_que_nao_explica():
    indicators_store.salvar_extracao("swa", "week", {
        "2026-09-06": {}, "2026-09-07": {},   # domingo e segunda: certas
        "2026-09-10": {}})                     # quinta, mas 09/10 trocado nao e semana
    assert reparo_semanas.reparar() == 0
    assert sorted(indicators_store.get_indicators("swa")["week"]) == [
        "2026-09-06", "2026-09-07", "2026-09-10"]
